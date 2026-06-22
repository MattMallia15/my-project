#!/usr/bin/env python3
"""Convert portfolio/company_notes.csv to a styled .xlsx workbook."""
import csv
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parent.parent
SRC = ROOT / "portfolio" / "company_notes.csv"
DST = ROOT / "portfolio" / "company_notes.xlsx"


def main() -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Company Notes"

    with SRC.open(newline="", encoding="utf-8") as f:
        rows = list(csv.reader(f))

    header, *data = rows

    header_font = Font(bold=True, color="FFFFFF", size=11)
    header_fill = PatternFill("solid", fgColor="1F3864")
    body_font = Font(size=10)
    thin = Side(style="thin", color="BFBFBF")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    wrap = Alignment(wrap_text=True, vertical="top")

    ws.append(header)
    for col_idx, _ in enumerate(header, 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(
            wrap_text=True, vertical="center", horizontal="center"
        )
        cell.border = border

    for row in data:
        ws.append(row)

    for row_idx in range(2, ws.max_row + 1):
        for col_idx in range(1, ws.max_column + 1):
            cell = ws.cell(row=row_idx, column=col_idx)
            cell.font = body_font
            cell.alignment = wrap
            cell.border = border
        if row_idx % 2 == 0:
            for col_idx in range(1, ws.max_column + 1):
                ws.cell(row=row_idx, column=col_idx).fill = PatternFill(
                    "solid", fgColor="F2F2F2"
                )

    widths = {
        "date": 12,
        "ticker": 16,
        "topic": 28,
        "key_facts": 55,
        "view_or_decision": 50,
        "next_step": 40,
        "sources": 22,
    }
    for col_idx, name in enumerate(header, 1):
        ws.column_dimensions[get_column_letter(col_idx)].width = widths.get(name, 20)

    ws.row_dimensions[1].height = 28
    for row_idx in range(2, ws.max_row + 1):
        ws.row_dimensions[row_idx].height = 90

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    wb.save(DST)
    print(f"Wrote {DST.relative_to(ROOT)} ({ws.max_row - 1} rows)")


if __name__ == "__main__":
    main()
